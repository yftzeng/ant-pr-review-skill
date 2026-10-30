"""Markdown first, then PDF and DOCX in Python, in the run's language (D-0007)."""
import re
import sys

import pytest

from ant_pr_review import export
from ant_pr_review.i18n import T
from ant_pr_review.jsonio import write_json

MD = """# 危機應變工作單

第一段，含**粗體**與待填〔　　〕。

## 立即行動

- [ ] 成立危機小組
- [x] 暫停排程貼文
1. 統一口徑
2. 指定發言人

| 項目 | 負責 | 截止 |
|---|---|---|
| 受影響範圍 | 品保主管 | 18:00 |

> 本文件尚未經過完整審查

---

??? a line the subset does not know
"""
WORDS = ["危機應變工作單", "第一段，含粗體與待填〔　　〕。", "立即行動", "成立危機小組", "暫停排程貼文",
         "統一口徑", "指定發言人", "受影響範圍", "品保主管", "18:00", "本文件尚未經過完整審查",
         "??? a line the subset does not know"]


def flat(s):
    return re.sub(r"[\s\x00]+", "", s)  # pypdf puts NULs between CJK glyphs; poppler does not


def pdf_text(path):
    from pypdf import PdfReader
    r = PdfReader(str(path))
    return r, flat("".join(p.extract_text() for p in r.pages))


def docx_text(path):
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for tb in d.tables for row in tb.rows for c in row.cells]
    parts += [p.text for s in d.sections for p in s.header.paragraphs]
    return d, flat("\n".join(parts))


def font_or_skip(lang):
    f = export.find_font(lang)
    if f is None:
        pytest.skip("no CJK font on this machine")
    return f


# @S-export-001
def test_subset_to_pdf_and_docx(tmp_path):
    font_or_skip("zh-TW")
    md = tmp_path / "sheet.md"
    md.write_text(MD, encoding="utf-8")
    res = export.export_md(md, "zh-TW", version="v1", when="2026-10-30 10:00")
    assert res["pdf"] == tmp_path / "sheet.pdf" and res["docx"] == tmp_path / "sheet.docx"
    r, ptext = pdf_text(res["pdf"])
    box = r.pages[0].mediabox
    assert (round(float(box.width)), round(float(box.height))) == (595, 842)  # A4
    _, dtext = docx_text(res["docx"])
    for w in WORDS:
        assert flat(w) in ptext, w
        assert flat(w) in dtext, w
    for text in (ptext, dtext):
        assert "□成立危機小組" in text and "■暫停排程貼文" in text
        assert "**" not in text and "[ ]" not in text
    kinds = [b[0] for b in export.parse_md(MD)]
    assert kinds == ["h", "p", "h", "li", "li", "li", "li", "table", "quote", "hr", "p"]


# @S-export-002
@pytest.mark.parametrize("lang,face", [("zh-TW", "tc"), ("zh-CN", "sc"), ("en", "")])
def test_language_picks_labels_and_font(tmp_path, lang, face):
    font_or_skip(lang)
    md = tmp_path / "s.md"
    md.write_text("# 標題 Title\n\n內容 body\n", encoding="utf-8")
    res = export.export_md(md, lang, version="v2", when="2026-10-30 10:00")
    t = T(lang)
    r, ptext = pdf_text(res["pdf"])
    for label in (t("export.confidential"), t("export.version"), t("export.time")):
        assert flat(label) in ptext, label
    assert "v2" in ptext and "2026-10-3010:00" in ptext
    assert flat(t("export.page", n=1, total=1)) in ptext
    fonts = [str(f.get_object()["/BaseFont"]).lower()
             for f in r.pages[0]["/Resources"]["/Font"].values()]
    assert fonts and all("cjk" in f or "han" in f or "hei" in f for f in fonts), fonts
    if face:
        assert all(face in f for f in fonts), fonts
    _, dtext = docx_text(res["docx"])
    assert flat(t("export.confidential")) in dtext and "v2" in dtext
    if lang != "en":
        assert t("export.confidential") != T("en")("export.confidential")


# @S-export-003
def test_missing_packages_keep_the_markdown(tmp_path, monkeypatch):
    md = tmp_path / "s.md"
    md.write_text(MD, encoding="utf-8")
    monkeypatch.setitem(sys.modules, "fpdf", None)
    monkeypatch.setitem(sys.modules, "docx", None)
    res = export.export_md(md, "zh-TW")
    assert res["pdf"] is None and res["docx"] is None
    assert res["reasons"] == {"pdf": "packages", "docx": "packages"}
    assert "pip install" in res["install"] and "[export]" in res["install"]
    assert md.read_text(encoding="utf-8") == MD
    assert not (tmp_path / "s.pdf").exists() and not (tmp_path / "s.docx").exists()


# @S-export-003 no font: no PDF, still a DOCX; @N-export-002 an error never raises
def test_no_font_still_writes_docx(tmp_path, monkeypatch):
    md = tmp_path / "s.md"
    md.write_text(MD, encoding="utf-8")
    monkeypatch.delenv("ANT_PR_REVIEW_FONT", raising=False)
    monkeypatch.setattr(export, "CANDIDATES", [])
    res = export.export_md(md, "zh-TW")
    assert res["pdf"] is None and res["reasons"]["pdf"] == "font"
    assert res["docx"] and res["docx"].exists() and "pdf" not in res.get("errors", {})
    assert not (tmp_path / "s.pdf").exists()
    def boom(blocks, out, *a, **k):
        out.write_text("half", encoding="utf-8")
        raise RuntimeError("disk full")
    monkeypatch.setattr(export, "_write_docx", boom)
    res = export.export_md(md, "zh-TW")
    assert res["docx"] is None and res["reasons"]["docx"] == "error"
    assert "disk full" in res["errors"]["docx"]
    assert not (tmp_path / "s.docx").exists() and md.exists()
    assert not list(tmp_path.glob("*.part"))


# @S-export-004
def test_font_found_by_name(tmp_path, monkeypatch):
    tc = font_or_skip("zh-TW")
    sc = export.find_font("zh-CN")
    assert "TC" in tc["family"] and "SC" in sc["family"]
    if tc["path"] == sc["path"]:  # one collection: the face is picked by name
        assert tc["index"] != sc["index"]
    mine = tmp_path / "mine.ttc"
    mine.symlink_to(tc["path"])
    monkeypatch.setenv("ANT_PR_REVIEW_FONT", str(mine))
    got = export.find_font("zh-TW")
    assert got["path"] == str(mine) and "TC" in got["family"]
    monkeypatch.setenv("ANT_PR_REVIEW_FONT", str(tmp_path / "missing.ttf"))
    assert export.find_font("zh-TW")["path"] == tc["path"]


def _run(tmp_path, flag):
    from ant_pr_review.workspace import create_run
    run = create_run(tmp_path, "x", models=["claude:opus"], backend="fake", lang="zh-TW")
    write_json(run.path("00-intake.json"), {"purpose": "drafting", "lang": "zh-TW",
                                            **({"export": True} if flag else {})})
    return run


# @S-export-005 @N-export-001
def test_export_flag_writes_beside_the_final(tmp_path):
    font_or_skip("zh-TW")
    from ant_pr_review.intake import build_intake_from_args
    assert build_intake_from_args("outbound", ["x"], lang="en", export=True)["export"] is True
    assert "export" not in build_intake_from_args("outbound", ["x"], lang="en")
    run = _run(tmp_path / "a", True)
    text = "我們自 11 月 1 日起免費更換。\n"
    run.path("10-final.md").write_text(text, encoding="utf-8")
    res = export.export_if_asked(run, "10-final.md")
    assert run.path("10-final.pdf").exists() and run.path("10-final.docx").exists()
    assert run.path("10-final.md").read_text(encoding="utf-8") == text
    t = T("zh-TW")
    _, ptext = pdf_text(res["pdf"])
    assert flat(t("export.signoff")) in ptext and flat(t("export.signoff.role")) in ptext
    export.drop_exports(run, "10-final.md")
    assert not run.path("10-final.pdf").exists() and not run.path("10-final.docx").exists()
    off = _run(tmp_path / "b", False)
    off.path("final.md").write_text(text, encoding="utf-8")
    assert export.export_if_asked(off, "final.md") is None
    assert not off.path("final.pdf").exists()


# @S-export-005 the report and decide call it; a final that is no longer clean takes its files along
def test_report_and_decide_export(tmp_path, monkeypatch):
    from ant_pr_review import decision, report
    calls = []
    monkeypatch.setattr(export, "export_if_asked", lambda run, rel: calls.append(rel))
    from test_drafting import drafting_run
    run = drafting_run(tmp_path)
    write_json(run.path("08-drafts/summary.json"), {"status": "done", "chosen": "V1"})
    run.path("08-drafts/V1.md").write_text("We replace it from 1 November.\n", encoding="utf-8")
    assert report.export_final(run)["written"] is True and calls == ["10-final.md"]
    run.path("10-final.pdf").write_text("old", encoding="utf-8")
    run.path("08-drafts/V1.md").write_text("TODO: date\n", encoding="utf-8")
    assert report.export_final(run)["written"] is False
    assert not run.path("10-final.pdf").exists()
    run.path("09-plans").mkdir(exist_ok=True)
    run.path("09-plans/A.md").write_text("A text\n", encoding="utf-8")
    decision.write_final(run, "A")
    assert calls[-1] == "final.md"
